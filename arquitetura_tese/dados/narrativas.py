"""Narrativas sintéticas vinculadas à base tabular — decisão D7(b).

Gera Boletins de Ocorrência / históricos forenses simulados, condicionados às
features de cada episódio de base_sintetica_risco.xlsx. Pessoas aparecem como
placeholders derivados dos ids sintéticos — privacidade por construção — e a
ligação pessoa↔episódios é preservada (habilita o grafo de reincidência, D6).

Cada documento emitido traz:
  - narrativa: texto simulando BO forense (âncoras lexicais para o BM25 e
    contexto denso para os encoders — Fase 2);
  - conceitos: rótulos de conceito para o dataset do probe CAV (D3);
  - dispositivos_sugeridos: rótulos fracos de relevância legal — bootstrap do
    gold set (§6.1); o especialista valida e corrige na anotação.
"""

import random
from pathlib import Path

BASE_XLSX = Path(__file__).resolve().parents[2] / "dados" / "base_sintetica_risco.xlsx"

ABERTURA_GENERICA = [
    "BOLETIM DE OCORRÊNCIA — Comparece à delegacia a vítima {vitima}, {idade_v} anos, "
    "para relatar fato ocorrido em {local}, envolvendo o requerido {agressor}, {idade_r} anos.",
    "TERMO DE DEPOIMENTO — A declarante {vitima}, {idade_v} anos, narra episódio ocorrido "
    "em {local}, atribuído ao requerido {agressor}, {idade_r} anos.",
    "HISTÓRICO FORENSE — Registro de episódio em {local}: vítima {vitima} ({idade_v} anos), "
    "requerido {agressor} ({idade_r} anos).",
]

ABERTURA_MPU = [
    "PEDIDO DE MEDIDA PROTETIVA — A requerente {vitima}, {idade_v} anos, requer proteção "
    "de urgência contra {agressor}, {idade_r} anos, com fundamento na Lei 11.340/2006, arts. 22 e 23.",
    "EXPEDIENTE DE URGÊNCIA — {vitima} solicita medidas protetivas em face de {agressor}, "
    "em contexto de violência doméstica e familiar contra a mulher (Lei 11.340/2006).",
]

CONDICIONAIS: list[tuple[str, list[str]]] = [
    ("descumpriu_medida", [
        "Consta que o requerido descumpriu medida protetiva de urgência vigente, "
        "aproximando-se da residência da vítima, nos termos do art. 24-A da Lei 11.340/2006.",
        "A vítima possui medida protetiva deferida; o agressor foi visto nas imediações "
        "do local de trabalho dela, em descumprimento à ordem judicial.",
        "Há registro de descumprimento de medida protetiva: o requerido enviou mensagens "
        "intimidadoras à vítima mesmo após a proibição judicial de contato.",
    ]),
    ("ameaca_morte", [
        "Relata a vítima que o agressor afirmou que a mataria caso procurasse a polícia novamente.",
        "O requerido teria dito à vítima, em tom intimidador, que ela 'não viveria para "
        "contar a história'.",
        "Constam ameaças de morte direcionadas à vítima e, por extensão, aos familiares que a abrigam.",
    ]),
    ("acesso_arma", [
        "Informa-se que o agressor possui arma de fogo em casa e já a exibiu para intimidar a vítima.",
        "A vítima relata que o requerido menciona a posse de arma de fogo durante as discussões.",
        "Há indício de acesso a arma de fogo pelo requerido, fator que agrava o risco de escalada.",
    ]),
    ("uso_alcool_drogas", [
        "O fato ocorreu enquanto o agressor apresentava visíveis sinais de embriaguez.",
        "Segundo a vítima, as agressões costumam ocorrer quando o requerido faz uso de bebida alcoólica.",
        "Consta uso habitual de álcool e/ou entorpecentes pelo agressor no contexto dos episódios.",
    ]),
    ("separacao_recente", [
        "O episódio ocorre poucas semanas após a separação do casal, período reconhecidamente "
        "associado a maior risco de violência.",
        "A vítima comunicou o fim do relacionamento dias antes do fato, o que teria motivado "
        "a reação do agressor.",
        "Consta que a separação recente antecedeu a escalada de ameaças narrada pela vítima.",
    ]),
    ("controle_coercitivo", [
        "Descreve-se padrão de controle coercitivo: monitoramento do celular, restrição de "
        "saídas e isolamento de familiares.",
        "A vítima relata controle financeiro e impedimento de trabalhar fora, em dinâmica de "
        "subordinação prolongada.",
        "Consta vigilância constante sobre a vítima, com exigência de prestação de contas "
        "de rotina e de amizades.",
    ]),
    ("filhos_comum", [
        "O casal possui filhos em comum, o que mantém contato frequente mesmo após a ruptura.",
        "Registra-se a existência de filhos comuns, usados pelo requerido como pretexto de aproximação.",
    ]),
    ("desemprego_agressor", [
        "O agressor encontra-se desempregado e passa a maior parte do tempo na residência.",
        "Consta situação de desemprego do requerido, apontada como fator de tensão no núcleo familiar.",
    ]),
]

ESCALADA = [
    "É a {n}ª ocorrência registrada envolvendo o mesmo agressor, indicando padrão de escalada.",
    "O histórico aponta {n} episódios anteriores atribuídos ao mesmo requerido, em progressão de gravidade.",
    "Registra-se reiteração: {n} fatos anteriores já documentados contra o mesmo agressor.",
]

ENCERRAMENTO = [
    "Diante do exposto, lavrou-se o presente termo e a vítima foi orientada quanto aos "
    "direitos previstos na Lei 11.340/2006.",
    "A ocorrência foi registrada e encaminhada à autoridade competente para as providências legais.",
    "Encerra-se o registro com orientação sobre medidas protetivas e encaminhamento à rede de apoio.",
]

# feature -> conceito para o dataset do probe CAV (D3/§4.6)
CONCEITOS = {
    "controle_coercitivo": "coercao",
    "descumpriu_medida": "descumprimento_medida",
    "ameaca_morte": "ameaca_morte",
    "acesso_arma": "arma_fogo",
    "uso_alcool_drogas": "abuso_substancias",
    "separacao_recente": "separacao_recente",
}

DISPOSITIVOS_BASE = [
    {"doc_id": "lei11340_art5", "lei": "Lei 11.340/2006", "artigo": "Art. 5º"},
    {"doc_id": "lei11340_art7", "lei": "Lei 11.340/2006", "artigo": "Art. 7º"},
]

DISPOSITIVOS_FEATURE = {
    "descumpriu_medida": {"doc_id": "lei11340_art24a", "lei": "Lei 11.340/2006", "artigo": "Art. 24-A"},
    "ameaca_morte": {"doc_id": "cp_art147", "lei": "Código Penal", "artigo": "Art. 147"},
    "controle_coercitivo": {"doc_id": "cp_art147a", "lei": "Código Penal", "artigo": "Art. 147-A"},
}

DISPOSITIVOS_CLASSE = {
    "medidas protetivas": {"doc_id": "lei11340_art19", "lei": "Lei 11.340/2006", "artigo": "Art. 19"},
    "inquérito": {"doc_id": "lei11340_art12", "lei": "Lei 11.340/2006", "artigo": "Art. 12"},
}


def _flag(v) -> bool:
    return str(v).strip() in {"1", "1.0", "True", "true"}


def _int(v, default: int = 0) -> int:
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return default


def placeholder(pessoa: str, id_sintetico) -> str:
    return f"{pessoa}_{str(id_sintetico)[:6].upper()}"


def gerar_narrativa(ep: dict, rng: random.Random) -> str:
    vitima = placeholder("VÍTIMA", ep["id_vitima_sintetica"])
    agressor = placeholder("AGRESSOR", ep["id_agressor_sintetico"])
    classe = str(ep.get("classe") or "").lower()
    ctx = {
        "vitima": vitima,
        "agressor": agressor,
        "idade_v": _int(ep.get("idade_vitima")),
        "idade_r": _int(ep.get("idade_requerido")),
        "local": ep.get("local") or "não informado",
    }
    banco = ABERTURA_MPU if "medida" in classe else ABERTURA_GENERICA
    partes = [rng.choice(banco).format(**ctx)]
    for feature, trechos in CONDICIONAIS:
        if _flag(ep.get(feature)):
            partes.append(rng.choice(trechos).format(**ctx))
    n_ant = _int(ep.get("n_episodios_anteriores"))
    if n_ant >= 1:
        partes.append(rng.choice(ESCALADA).format(n=n_ant + 1, **ctx))
    partes.append(rng.choice(ENCERRAMENTO))
    return " ".join(partes)


def conceitos(ep: dict) -> list[str]:
    rotulos = [c for f, c in CONCEITOS.items() if _flag(ep.get(f))]
    if _int(ep.get("n_episodios_anteriores")) >= 2:
        rotulos.append("escalada")
    return rotulos


def dispositivos_sugeridos(ep: dict) -> list[dict]:
    vistos = {d["doc_id"]: d for d in DISPOSITIVOS_BASE}
    for f, d in DISPOSITIVOS_FEATURE.items():
        if _flag(ep.get(f)):
            vistos[d["doc_id"]] = d
    classe = str(ep.get("classe") or "").lower()
    for chave, d in DISPOSITIVOS_CLASSE.items():
        if chave in classe:
            vistos[d["doc_id"]] = d
    return list(vistos.values())


def iterar_episodios(xlsx_path: str | Path):
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    ws = wb.active
    linhas = ws.iter_rows(values_only=True)
    header = next(linhas)
    for row in linhas:
        yield dict(zip(header, row))
    wb.close()


def gerar_documento(ep: dict, rng: random.Random) -> dict:
    return {
        "id_episodio": ep["id_episodio"],
        "id_agressor": ep["id_agressor_sintetico"],
        "id_vitima": ep["id_vitima_sintetica"],
        "classe": ep.get("classe"),
        "local": ep.get("local"),
        "narrativa": gerar_narrativa(ep, rng),
        "conceitos": conceitos(ep),
        "dispositivos_sugeridos": dispositivos_sugeridos(ep),
    }
