"""Converte Markdown leve para .docx (python-docx).

Suporta: #..#### títulos, **negrito**, `código`, tabelas pipe, listas
(-/*/1.), blocos ```, blockquote, ---, links [t](u) -> t (u).
Uso: python scripts/md_para_docx.py entrada.md saida.docx
"""
import re
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def _runs(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            par.add_run(f"{m.group(1)} ({m.group(2)})" if m else part)
        else:
            par.add_run(part)


def converter(md_path, docx_path, titulo=None):
    lines = open(md_path, encoding="utf-8").read().splitlines()
    doc = Document()
    st = doc.styles["Normal"].font
    st.name, st.size = "Calibri", Pt(11)
    if titulo:
        h = doc.add_heading(titulo, 0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    i, in_code, in_table = 0, False, False
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()

        if s.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            r = doc.add_paragraph().add_run(ln)
            r.font.name, r.font.size = "Consolas", Pt(9)
            i += 1
            continue

        if s.startswith("|") and s.endswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                t = doc.add_table(len(rows), max(len(r) for r in rows))
                t.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, c in enumerate(row):
                        cell = t.cell(ri, ci)
                        cell.text = ""
                        _runs(cell.paragraphs[0], c)
                        for r in cell.paragraphs[0].runs:
                            r.font.size = Pt(9)
                            if ri == 0:
                                r.bold = True
            continue

        if m := re.match(r"^(#{1,4})\s+(.*)", s):
            if titulo and len(m.group(1)) == 1:
                titulo = None  # pula H1 próprio do md (já usado como título)
                i += 1
                continue
            doc.add_heading(re.sub(r"\*\*", "", m.group(2)),
                            level=min(len(m.group(1)), 4))
        elif s == "---":
            pass
        elif s.startswith(">"):
            p = doc.add_paragraph()
            r = p.add_run(s.lstrip("> "))
            r.italic = True
        elif re.match(r"^[-*]\s+", s):
            _runs(doc.add_paragraph(style="List Bullet"), s[2:])
        elif re.match(r"^\d+[.)]\s+", s):
            _runs(doc.add_paragraph(style="List Number"),
                  re.sub(r"^\d+[.)]\s+", "", s))
        elif s:
            _runs(doc.add_paragraph(), s)
        i += 1
    doc.save(docx_path)
    print("->", docx_path)


if __name__ == "__main__":
    converter(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)
