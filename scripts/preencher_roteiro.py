"""Preenche 'Roteiro Metodologia.docx' com tabela do dicionário, amostra de
narrativa e figuras do projeto. Salva cópia '- preenchido.docx'."""
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

RAIZ = r"D:\experimento_silvia\judicata"
SRC = r"D:\experimento_silvia\Roteiro Metodologia.docx"
DST = r"D:\experimento_silvia\Roteiro Metodologia - preenchido.docx"
FIG = RAIZ + r"\dados\figuras\\"

# --- dicionário de dados (lido do md) ---------------------------------------
md = open(RAIZ + r"\dados\dicionario_base.md", encoding="utf-8").read()
linhas = [l for l in md.splitlines() if l.startswith("| `")]
dicio = [[c.strip().strip("`") for c in l.strip("|").split("|")] for l in linhas]
narr = md.split("## Narrativa sintética renderizada (exemplo)")[1]
narr = narr.split("```")[1].strip()

doc = Document(SRC)
paras = doc.paragraphs


def _img(p, arquivo, largura=6.2):
    for r in list(p.runs):
        r.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(FIG + arquivo, width=Inches(largura))


def _img_antes(ref_p, arquivo, largura=6.2):
    p = ref_p.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(FIG + arquivo, width=Inches(largura))
    return p


def _legenda_antes(ref_p, texto):
    p = ref_p.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texto)
    r.font.size = Pt(10)
    return p


def _caixa(p, texto):
    for r in list(p.runs):
        r.text = ""
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.right_indent = Inches(0.4)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "444444")
        bdr.append(e)
    pPr.append(bdr)
    r = p.add_run(texto)
    r.font.size = Pt(10)
    r.italic = True


def _tabela_apos(p, cab, linhas_):
    t = doc.add_table(len(linhas_) + 1, len(cab))
    tblPr = t._tbl.tblPr
    bdr = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "000000")
        bdr.append(e)
    tblPr.append(bdr)
    for j, h in enumerate(cab):
        c = t.cell(0, j)
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9)
    for i, row in enumerate(linhas_, 1):
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            r = c.paragraphs[0].add_run(v)
            r.font.size = Pt(9)
    p._p.addnext(t._tbl)
    p.text = ""


for p in paras:
    s = p.text.strip()
    if s.startswith("[INSERIR TABELA"):
        _tabela_apos(p, ["coluna", "papel", "descrição", "r(y)"], dicio)
    elif s.startswith("[INSERIR TEXTO - AMOSTRA"):
        _caixa(p, narr)
    elif s.startswith("[INSERIR FIGURA 1"):
        _img(p, "fig1_roc_importancia.png")
    elif s.startswith("[INSERIR FIGURA 2"):
        _img(p, "fig2_hitrate.png")
    elif s.startswith("[INSERIR FIGURA 3"):
        _img(p, "fig3_latencia.png")
    elif s.startswith("[INSERIR FIGURA 5"):
        _img_antes(p, "fig10_matriz_confusao.png")
        _legenda_antes(p, "Figura 5.5a. Matriz de confusão no holdout "
                        "(n=3.993, limiar 0,50): VP=1.294, FP=524, "
                        "VN=1.463, FN=712.")
        _img(p, "fig5_calibracao.png")
    elif s.startswith("Figura 5.1."):
        _img_antes(p, "fig7_arquitetura.png")

doc.save(DST)
print("->", DST)
